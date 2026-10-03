import io
import logging
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from database import Database

logger = logging.getLogger(__name__)

HEADERS = ["Session ID", "Subject", "Student ID", "Student Name", "Date",
           "Entry Time", "Exit Time", "Duration (min)", "Status", "Recognition Method"]


class ExcelService:
    @staticmethod
    def build_attendance_workbook(filters=None):
        """Returns an in-memory .xlsx (BytesIO). filters: session_id / date / student_id (empty ones ignored)."""
        db = Database.get_db()
        if db is None:
            raise RuntimeError("Database connection not available")

        query = {k: v for k, v in (filters or {}).items() if v}
        records = list(db.attendance.find(query, {"_id": 0}).sort("created_at", -1))

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Attendance Log"
        ws.append(HEADERS)
        fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
        font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        for col in range(1, len(HEADERS) + 1):
            c = ws.cell(row=1, column=col)
            c.fill, c.font = fill, font
            c.alignment = Alignment(horizontal="center", vertical="center")

        for r in records:
            ws.append([
                r.get("session_id", ""), r.get("subject", ""), r.get("student_id", ""),
                r.get("student_name", ""), r.get("date", ""), r.get("entry_time", ""),
                r.get("exit_time") or "N/A", r.get("duration_minutes", 0),
                r.get("status", "Present"), r.get("recognition_method", ""),
            ])

        for col in ws.columns:
            width = max(len(str(c.value or "")) for c in col)
            ws.column_dimensions[get_column_letter(col[0].column)].width = max(width + 3, 12)
        ws.freeze_panes = "A2"

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return buf
